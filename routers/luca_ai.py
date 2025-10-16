# TESTE PARA VERIFICAR O GIT
# /routers/luca_ai.py
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from pydantic import BaseModel
from typing import List, Optional, Dict
import os
import json
import re
from datetime import datetime, date, timedelta
import requests
from pathlib import Path
import calendar
import random
import io

# --- NOVAS DEPENDÊNCIAS ---
# Certifique-se de instalar estas bibliotecas com:
# pip install PyMuPDF openpyxl python-docx
import fitz  # PyMuPDF
import openpyxl
import docx

from main_api import send_command_to_agent, execute_query_via_agent
from dependencies import get_company_fk, EmpresaInfo, verificar_empresa

router = APIRouter()
KNOWLEDGE_BASE_DIR = "base_conhecimento_local"

def load_prompts_from_file() -> Dict:
    """Carrega os templates de prompt do arquivo PROMPT.txt."""
    prompts = {}
    try:
        prompts_path = Path(KNOWLEDGE_BASE_DIR) / "PROMPT.txt"
        if not prompts_path.exists():
            print("ERRO CRÍTICO: O arquivo 'PROMPT.txt' não foi encontrado na base de conhecimento local.")
            return {}
        
        content = prompts_path.read_text(encoding='utf-8')
        
        pattern = re.compile(r"\[PROMPT:\s*([^\]]+)\]\n?(.*?)\n?\[ENDPROMPT\]", re.DOTALL)
        matches = pattern.findall(content)
        
        for name, text in matches:
            prompts[name.strip()] = text.strip()
            
        if not prompts:
            print("ERRO CRÍTICO: Nenhum prompt válido foi encontrado no arquivo 'PROMPT.txt'. Verifique a formatação.")
            
        return prompts
        
    except Exception as e:
        print(f"ERRO CRÍTICO ao carregar ou parsear o arquivo de prompts 'PROMPT.txt': {e}")
        return {}

PROMPTS = load_prompts_from_file()

def load_knowledge_from_local_files(local_path_str: str) -> str:
    local_path = Path(local_path_str)
    if not local_path.exists():
        return ""
    
    knowledge_content = []
    for file_path in local_path.iterdir():
        if file_path.name.lower() != "prompt.txt" and file_path.suffix.lower() == '.json':
            try:
                knowledge_content.append(file_path.read_text(encoding='utf-8'))
            except Exception as e:
                print(f"AVISO: Falha ao ler o arquivo de conhecimento '{file_path.name}': {e}")
            
    return "\n\n".join(knowledge_content)


class LucaRequest(BaseModel):
    prompt: str

class ReportData(BaseModel):
    title: str; summary: str; table_headers: List[str]; table_rows: List[dict]

class LucaResponse(BaseModel):
    answer: str
    report_data: Optional[ReportData] = None

def call_gemini_api(prompt: str, api_key: str):
    if not api_key:
        raise HTTPException(status_code=500, detail="A GEMINI_API_KEY não foi configurada no ambiente do servidor.")
    
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-pro-latest:generateContent?key={api_key}"
    
    headers = {'Content-Type': 'application/json'}
    payload = {"contents": [{"parts": [{"text": prompt}]}], "safetySettings": [{"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_NONE"}, {"category": "HARM_CATEGORY_HATE_SPEECH", "threshold": "BLOCK_NONE"}, {"category": "HARM_CATEGORY_SEXUALLY_EXPLICIT", "threshold": "BLOCK_NONE"}, {"category": "HARM_CATEGORY_DANGEROUS_CONTENT", "threshold": "BLOCK_NONE"}]}
    try:
        response = requests.post(url, headers=headers, json=payload, timeout=90)
        response.raise_for_status()
        data = response.json()
        if not data.get('candidates'):
             raise HTTPException(status_code=500, detail="A resposta da IA foi bloqueada ou veio vazia.")
        return data['candidates'][0]['content']['parts'][0]['text']
    except requests.exceptions.HTTPError as http_err:
        print(f"ERRO HTTP da API Gemini: {http_err.response.status_code} - {http_err.response.text}")
        raise HTTPException(status_code=http_err.response.status_code, detail=f"Erro na API do Google AI: {http_err.response.text}")
    except Exception as e:
        print(f"ERRO Inesperado na chamada da API: {e}")
        raise HTTPException(status_code=500, detail=f"Ocorreu um erro inesperado na API: {e}")

async def generate_goal_simulation_response(company_cnpj: str, id_empresa: str, api_key: str) -> str:
    today = date.today()
    current_year, current_month = today.year, today.month
    
    sql_meta = "SELECT VALOR FROM DBMETAS WHERE ID_EMPRESA = ? AND ANO = ? AND MES = ? AND INDICADOR = 'FATURAMENTO_MENSAL'"
    meta_res = await execute_query_via_agent(company_cnpj, sql_meta, [id_empresa, current_year, current_month])
    
    if not meta_res or not meta_res[0].get('VALOR'):
        return "Não encontrei uma meta de faturamento definida para este mês. Para que eu possa ajudar, por favor, cadastre uma meta no painel específico."

    valor_meta = float(meta_res[0]['VALOR'])

    sql_performance = "SELECT SUM(CAST(VALORLIQUIDO AS DOUBLE PRECISION)) AS TOTAL_VENDIDO, COUNT(CODIGO) AS NUM_PEDIDOS FROM TVENPEDIDO WHERE EMPRESA = ? AND STATUS = 'EFE' AND TIPOVENDA = 'NM' AND EXTRACT(YEAR FROM DATAEFE) = ? AND EXTRACT(MONTH FROM DATAEFE) = ?"
    perf_res = await execute_query_via_agent(company_cnpj, sql_performance, [id_empresa, current_year, current_month])
    
    total_vendido = float(perf_res[0].get('TOTAL_VENDIDO') or 0.0) if perf_res else 0.0
    num_pedidos = int(perf_res[0].get('NUM_PEDIDOS') or 0) if perf_res else 0

    valor_restante = valor_meta - total_vendido
    if valor_restante <= 0:
        return f"🏆 Parabéns! A meta de faturamento de **R$ {valor_meta:,.2f}** para este mês já foi atingida e superada em **R$ {abs(valor_restante):,.2f}**!"

    _, total_dias_mes = calendar.monthrange(current_year, current_month)
    dias_restantes = total_dias_mes - today.day + 1
    
    if dias_restantes <= 0:
        return f"O mês já terminou. A meta era de R$ {valor_meta:,.2f} e o faturamento foi de R$ {total_vendido:,.2f}. Faltaram R$ {valor_restante:,.2f} para atingir o objetivo."

    media_diaria_necessaria = valor_restante / dias_restantes if dias_restantes > 0 else 0
    ticket_medio_atual = total_vendido / num_pedidos if num_pedidos > 0 else 0
    vendas_diarias_atuais = num_pedidos / today.day if today.day > 0 else 0
    projecao_vendas_restantes = vendas_diarias_atuais * dias_restantes
    
    novo_ticket_medio_sugerido = (valor_restante / projecao_vendas_restantes) if projecao_vendas_restantes > 0 else 0
    novas_vendas_necessarias = (valor_restante / ticket_medio_atual) if ticket_medio_atual > 0 else 0

    prompt_template = PROMPTS.get("goal_simulation")
    if not prompt_template: return "Erro: Template de prompt 'goal_simulation' não encontrado."

    prompt_contexto = prompt_template.format(
        valor_meta=valor_meta,
        total_vendido=total_vendido,
        valor_restante=valor_restante,
        dias_restantes=dias_restantes,
        media_diaria_necessaria=media_diaria_necessaria,
        ticket_medio_atual=ticket_medio_atual,
        novo_ticket_medio_sugerido=novo_ticket_medio_sugerido,
        novas_vendas_necessarias=novas_vendas_necessarias
    )
    return call_gemini_api(prompt_contexto, api_key)

async def generate_promotion_ideas(company_cnpj: str, id_empresa: str, api_key: str) -> str:
    sql_bundles = "SELECT FIRST 5 g1.DESCRICAO as NOME_A, g2.DESCRICAO as NOME_B, COUNT(*) as VEZES_COMPRADOS_JUNTOS FROM TVENPRODUTO p1 JOIN TVENPRODUTO p2 ON p1.PEDIDO = p2.PEDIDO AND p1.PRODUTO < p2.PRODUTO AND p1.EMPRESA = p2.EMPRESA JOIN TVENPEDIDO ped ON p1.PEDIDO = ped.CODIGO AND p1.EMPRESA = ped.EMPRESA JOIN TESTPRODUTOGERAL g1 ON p1.PRODUTO = g1.CODIGO JOIN TESTPRODUTOGERAL g2 ON p2.PRODUTO = g2.CODIGO WHERE ped.EMPRESA = ? AND ped.STATUS = 'EFE' AND ped.TIPOVENDA = 'NM' AND ped.DATAEFE >= DATEADD(-90 DAY TO CURRENT_DATE) GROUP BY 1, 2 ORDER BY 3 DESC"
    bundles_res = await execute_query_via_agent(company_cnpj, sql_bundles, [id_empresa])

    sql_volume = "SELECT FIRST 5 g.DESCRICAO, SUM(CAST(i.QTDE AS DOUBLE PRECISION)) as TOTAL_QTDE FROM TVENPEDIDO p JOIN TVENPRODUTO i ON p.CODIGO = i.PEDIDO AND p.EMPRESA = i.EMPRESA JOIN TESTPRODUTOGERAL g ON i.PRODUTO = g.CODIGO WHERE p.EMPRESA = ? AND p.STATUS = 'EFE' AND p.TIPOVENDA = 'NM' AND p.DATAEFE >= DATEADD(-90 DAY TO CURRENT_DATE) GROUP BY 1 ORDER BY 2 DESC"
    volume_res = await execute_query_via_agent(company_cnpj, sql_volume, [id_empresa])

    if not bundles_res and not volume_res:
        return "💡 Para que eu possa sugerir promoções eficazes, preciso de um histórico maior de vendas. Continue registrando seus pedidos e em breve terei insights valiosos para você!"

    bundles_data_str = ""
    if bundles_res:
        bundles_data_str = f"**Produtos Frequentemente Comprados Juntos (últimos 90 dias):**\n`{json.dumps(bundles_res, indent=2, default=str)}`\n\n"

    volume_data_str = ""
    if volume_res:
        volume_data_str = f"**Produtos Mais Vendidos em Quantidade (últimos 90 dias):**\n`{json.dumps(volume_res, indent=2, default=str)}`\n\n"

    prompt_template = PROMPTS.get("promotion_ideas")
    if not prompt_template: return "Erro: Template de prompt 'promotion_ideas' não encontrado."

    prompt = prompt_template.format(bundles_data=bundles_data_str, volume_data=volume_data_str)
    
    return call_gemini_api(prompt, api_key)

async def generate_surprise_insight(company_cnpj: str, id_empresa: str, api_key: str) -> str:
    async def analyze_worst_selling_day():
        sql = "SELECT EXTRACT(WEEKDAY FROM DATAEFE) as DIA_SEMANA, AVG(CAST(VALORLIQUIDO AS DOUBLE PRECISION)) as MEDIA_VENDAS FROM TVENPEDIDO WHERE EMPRESA = ? AND STATUS = 'EFE' AND TIPOVENDA = 'NM' AND DATAEFE >= DATEADD(-90 DAY TO CURRENT_DATE) GROUP BY 1 ORDER BY 2 ASC"
        results = await execute_query_via_agent(company_cnpj, sql, [id_empresa])
        if not results: return None
        
        dias_semana = ["Domingo", "Segunda-feira", "Terça-feira", "Quarta-feira", "Quinta-feira", "Sexta-feira", "Sábado"]
        worst_day_data = results[0]
        worst_day_name = dias_semana[int(worst_day_data['DIA_SEMANA'])]
        
        prompt_template = PROMPTS.get("surprise_worst_day")
        if not prompt_template: return "Erro: Template 'surprise_worst_day' não encontrado."
        
        prompt_context = prompt_template.format(worst_day_name=worst_day_name)
        return await call_gemini_api(prompt_context, api_key)

    async def analyze_high_margin_low_volume():
        sql = "WITH ProductStats AS (SELECT p.PRODUTO, COALESCE(g.DESCRICAOREDUZIDA, g.DESCRICAO) as NOME, SUM(CAST(i.VLRLIQUIDO AS DOUBLE PRECISION)) as FATURAMENTO_TOTAL, SUM(CAST(i.VLRLIQUIDO AS DOUBLE PRECISION) - (CAST(i.QTDE AS DOUBLE PRECISION) * CAST(i.CUSTOFINAL AS DOUBLE PRECISION))) as LUCRO_TOTAL, SUM(CAST(i.QTDE AS DOUBLE PRECISION)) as VOLUME_TOTAL FROM TVENPEDIDO p JOIN TVENPRODUTO i ON p.CODIGO = i.PEDIDO AND p.EMPRESA = i.EMPRESA JOIN TESTPRODUTOGERAL g ON i.PRODUTO = g.CODIGO WHERE p.EMPRESA = ? AND p.STATUS = 'EFE' AND p.TIPOVENDA = 'NM' AND p.DATAEFE >= DATEADD(-180 DAY TO CURRENT_DATE) AND i.VLRLIQUIDO > 0 GROUP BY 1, 2), RankedProducts AS (SELECT NOME, (LUCRO_TOTAL / FATURAMENTO_TOTAL) * 100 as MARGEM, VOLUME_TOTAL, NTILE(5) OVER (ORDER BY (LUCRO_TOTAL / FATURAMENTO_TOTAL) DESC) as MARGIN_QUINTILE, NTILE(2) OVER (ORDER BY VOLUME_TOTAL ASC) as VOLUME_HALF FROM ProductStats WHERE FATURAMENTO_TOTAL > 0 AND VOLUME_TOTAL > 0) SELECT FIRST 3 NOME, MARGEM FROM RankedProducts WHERE MARGIN_QUINTILE = 1 AND VOLUME_HALF = 1 ORDER BY MARGEM DESC"
        results = await execute_query_via_agent(company_cnpj, sql, [id_empresa])
        if not results: return None

        prompt_template = PROMPTS.get("surprise_high_margin")
        if not prompt_template: return "Erro: Template 'surprise_high_margin' não encontrado."

        prompt_context = prompt_template.format(results=json.dumps(results, indent=2, default=str))
        return await call_gemini_api(prompt_context, api_key)

    possible_analyses = [analyze_worst_selling_day, analyze_high_margin_low_volume]
    chosen_analysis = random.choice(possible_analyses)
    
    try:
        result = await chosen_analysis()
        if result: return result
        return "🎲 Tentei uma análise surpresa, mas não encontrei dados suficientes para gerar um insight interessante no momento. Talvez com mais alguns dias de vendas eu encontre algo!"
    except Exception as e:
        print(f"ERRO em análise surpresa: {e}")
        return "Opa, tive um problema ao tentar fazer a análise surpresa. Por favor, tente novamente."

@router.get("/luca/history", response_model=Dict[str, List[Dict[str, str]]])
async def get_chat_history(empresa_info: EmpresaInfo = Depends(verificar_empresa)):
    try:
        user_id, company_id = empresa_info.uid, empresa_info.company_id
        parametros = {"user_id": user_id}
        history = await send_command_to_agent(company_id, "carregar_historico", parametros)
        return history
    except Exception as e:
        print(f"ERRO ao buscar histórico de chat via agente: {e}")
        raise HTTPException(status_code=500, detail="Não foi possível carregar o histórico de conversas do agente local.")

@router.post("/luca/chat", response_model=LucaResponse)
async def handle_luca_chat(request: LucaRequest, empresa_info: EmpresaInfo = Depends(verificar_empresa), id_empresa: str = Depends(get_company_fk)):
    try:
        if not PROMPTS:
            raise HTTPException(status_code=500, detail="O arquivo de prompts (PROMPT.txt) não foi carregado ou está vazio. Verifique os logs do servidor.")

        # ### MELHORIA APLICADA AQUI ###
        # A chave de API agora é lida a partir das variáveis de ambiente do servidor.
        # Configure a variável 'GEMINI_API_KEY' no seu ambiente de produção (Render).
        api_key = os.environ.get("GEMINI_API_KEY")
        
        user_id, company_id = empresa_info.uid, empresa_info.company_id
        current_date_str = datetime.now().strftime('%Y-%m-%d')
        
        full_history = await send_command_to_agent(company_id, "carregar_historico", {"user_id": user_id})
        today_history = full_history.get(current_date_str, [])
        new_messages = [{"role": "user", "content": request.prompt}]
        
        prompt_lower = request.prompt.lower()
        goal_keywords = ["atingir a meta", "alcançar a meta", "bater a meta"]
        promotion_keywords = ["promoção", "promocao", "oferta", "desconto", "campanha"]
        surprise_keyword = "me surpreenda!"

        final_answer = None
        if surprise_keyword in prompt_lower:
            final_answer = await generate_surprise_insight(company_id, id_empresa, api_key)
        elif any(keyword in prompt_lower for keyword in promotion_keywords):
            final_answer = await generate_promotion_ideas(company_id, id_empresa, api_key)
        elif any(keyword in prompt_lower for keyword in goal_keywords):
            final_answer = await generate_goal_simulation_response(company_id, id_empresa, api_key)

        if final_answer:
            new_messages.append({"role": "model", "content": final_answer})
            await send_command_to_agent(company_id, "salvar_historico", {"user_id": user_id, "date_str": current_date_str, "new_messages": new_messages})
            return LucaResponse(answer=final_answer)

        knowledge_base_content = load_knowledge_from_local_files(KNOWLEDGE_BASE_DIR)
        formatted_history = "\n".join([f"  - {msg['role']}: {msg['content']}" for msg in today_history[-8:]])

        prompt_template = PROMPTS.get("sql_generation")
        if not prompt_template: return LucaResponse(answer="Erro: Template de prompt 'sql_generation' não encontrado.")

        full_prompt_for_sql = prompt_template.format(
            formatted_history=formatted_history,
            knowledge_base_content=knowledge_base_content,
            id_empresa=id_empresa,
            current_date_str=current_date_str,
            prompt=request.prompt
        )

        generated_text = call_gemini_api(full_prompt_for_sql, api_key)
        
        generated_sql = None
        if "```sql" in generated_text:
            parts = generated_text.split("```sql", 1)
            if len(parts) > 1: generated_sql = parts[1].split("```")[0].strip()

        if not generated_sql or not generated_sql.upper().startswith("SELECT"):
            new_messages.append({"role": "model", "content": generated_text})
            await send_command_to_agent(company_id, "salvar_historico", {"user_id": user_id, "date_str": current_date_str, "new_messages": new_messages})
            return LucaResponse(answer=generated_text)

        print(f"INFO: SQL Gerado e extraído: {generated_sql}")
        
        query_result = await execute_query_via_agent(company_id, generated_sql, [])

        if query_result is None or not query_result:
            return LucaResponse(answer="Realizei a consulta, mas não encontrei nenhum resultado para sua pergunta.")

        prompt_template_summary = PROMPTS.get("summary_generation")
        if not prompt_template_summary: return LucaResponse(answer="Erro: Template 'summary_generation' não encontrado.")
        
        prompt_for_summary = prompt_template_summary.format(
            prompt=request.prompt,
            query_result=json.dumps(query_result, indent=2, default=str)
        )
        final_answer = call_gemini_api(prompt_for_summary, api_key)
        
        new_messages.append({"role": "model", "content": final_answer})
        await send_command_to_agent(company_id, "salvar_historico", {"user_id": user_id, "date_str": current_date_str, "new_messages": new_messages})

        report_data = None
        if any(keyword in request.prompt.lower() for keyword in ["relatório", "tabela", "liste"]):
            report_data = ReportData(
                title=f"Relatório para: {request.prompt}", summary=final_answer,
                table_headers=list(query_result[0].keys()) if query_result else [],
                table_rows=query_result
            )

        return LucaResponse(answer=final_answer, report_data=report_data)

    except Exception as e:
        if isinstance(e, HTTPException): raise e
        raise HTTPException(status_code=500, detail=f"Ocorreu um erro inesperado no servidor do LUCA: {str(e)}")

@router.post("/luca/upload-and-analyze", response_model=LucaResponse)
async def handle_file_upload(
    empresa_info: EmpresaInfo = Depends(verificar_empresa), 
    file: UploadFile = File(...)
):
    """
    Recebe um arquivo (PDF, XLSX, DOCX, TXT, etc.), extrai seu conteúdo textual e usa a IA para gerar um resumo.
    """
    try:
        if not file.filename:
            raise HTTPException(status_code=400, detail="Nome do arquivo não fornecido.")
            
        allowed_content_types = [
            "text/plain", "text/csv", "application/json",
            "application/pdf",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", # XLSX
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"  # DOCX
        ]
        if file.content_type not in allowed_content_types:
            raise HTTPException(status_code=400, detail=f"Tipo de arquivo não suportado: {file.content_type}. Use .txt, .csv, .json, .pdf, .xlsx ou .docx.")

        max_size = 5 * 1024 * 1024
        contents_bytes = await file.read()
        if len(contents_bytes) > max_size:
            raise HTTPException(status_code=413, detail="Arquivo muito grande. O limite é de 5MB.")

        file_content = ""
        try:
            if file.content_type == "application/pdf":
                with fitz.open(stream=contents_bytes, filetype="pdf") as doc:
                    file_content = "".join(page.get_text() for page in doc)
            
            elif file.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                workbook = openpyxl.load_workbook(filename=io.BytesIO(contents_bytes))
                text_parts = []
                for sheet in workbook.worksheets:
                    text_parts.append(f"--- Planilha: {sheet.title} ---\n")
                    for row in sheet.iter_rows():
                        row_text = [str(cell.value) for cell in row if cell.value is not None]
                        if row_text:
                            text_parts.append("\t".join(row_text) + "\n")
                file_content = "".join(text_parts)

            elif file.content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                doc = docx.Document(io.BytesIO(contents_bytes))
                file_content = "\n".join([para.text for para in doc.paragraphs])
            
            else:
                encodings_to_try = ['utf-8', 'latin-1', 'windows-1252']
                for encoding in encodings_to_try:
                    try:
                        file_content = contents_bytes.decode(encoding)
                        print(f"INFO: Arquivo de texto decodificado com sucesso usando {encoding}.")
                        break
                    except UnicodeDecodeError:
                        continue
                if not file_content:
                     raise ValueError("Não foi possível decodificar o arquivo de texto.")

        except Exception as e:
            print(f"ERRO ao extrair conteúdo do arquivo {file.filename}: {e}")
            raise HTTPException(status_code=400, detail=f"Falha ao processar o conteúdo do arquivo. Ele pode estar corrompido ou em um formato não suportado. Erro: {e}")
        
        if not file_content.strip():
            raise HTTPException(status_code=400, detail="O arquivo parece estar vazio ou não contém texto extraível.")

        # ### MELHORIA APLICADA AQUI ###
        # A chave de API também é lida do ambiente para esta função.
        api_key = os.environ.get("GEMINI_API_KEY")
        
        prompt_template = PROMPTS.get("file_summary")
        if not prompt_template:
            raise HTTPException(status_code=500, detail="Template de prompt 'file_summary' não encontrado no servidor.")
            
        full_prompt = prompt_template.format(
            filename=file.filename,
            file_content=file_content
        )

        final_answer = call_gemini_api(full_prompt, api_key)
        
        user_id, company_id = empresa_info.uid, empresa_info.company_id
        current_date_str = datetime.now().strftime('%Y-%m-%d')
        new_messages = [
            {"role": "user", "content": f"Enviou o arquivo: **{file.filename}**"},
            {"role": "model", "content": final_answer}
        ]
        await send_command_to_agent(company_id, "salvar_historico", {"user_id": user_id, "date_str": current_date_str, "new_messages": new_messages})

        return LucaResponse(answer=final_answer)

    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        print(f"ERRO Inesperado no upload de arquivo: {e}")
        raise HTTPException(status_code=500, detail=f"Ocorreu um erro inesperado no servidor ao processar o arquivo: {str(e)}")